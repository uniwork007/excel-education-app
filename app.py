import os
import re
import unicodedata
import openpyxl
import requests
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

load_dotenv()

app = Flask(__name__)
app.secret_key = 'super-secret-key-change-this'

SUPABASE_URL = os.getenv('SUPABASE_URL')
SUPABASE_KEY = os.getenv('SUPABASE_KEY')

HEADERS = {
    "apikey": SUPABASE_KEY,
    "authorization": f"Bearer {SUPABASE_KEY}",
    "content-type": "application/json",
    "prefer": "return=representation"
}

ALLOWED_EXTENSIONS = {'xlsx', 'xls'}
UPLOAD_FOLDER = 'uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def is_date_cell(cell):
    return cell.is_date or cell.number_format in ['yyyy/mm/dd', 'yyyy-mm-dd', 'yyyy"年"m"月"d"日"']

def analyze_excel_file(filepath, stage_id):
    errors = []
    warnings = []
    cell_details = []
    
    try:
        wb = openpyxl.load_workbook(filepath, data_only=False)
        ws = wb.active
        
        url = f"{SUPABASE_URL}/answer_keys?stage_id=eq.{stage_id}"
        resp = requests.get(url, headers=HEADERS)
        expected_cells = resp.json() if resp.status_code == 200 else []
        
        for exp in expected_cells:
            cell_ref = exp['cell_ref']
            expected_val = str(exp['expected_value']).strip() if exp.get('expected_value') else None
            expected_type = exp.get('expected_type')
            
            cell = ws[cell_ref]
            actual_val = str(cell.value).strip() if cell.value is not None else ""
            
            norm_actual = unicodedata.normalize('NFKC', actual_val)
            norm_expected = unicodedata.normalize('NFKC', expected_val) if expected_val else ""
            
            detail = {
                "cell": cell_ref,
                "value": actual_val,
                "status": "OK",
                "msg": ""
            }
            
            # 全角文字が含まれているかチェック
            looks_like_formula = actual_val.startswith("=") or any(fn in actual_val.upper() for fn in ["SUM", "AVERAGE", "COUNT", "IF", "VLOOKUP", "XLOOKUP"])
            if looks_like_formula:
                has_zenkaku = any(ord(c) > 0x7F for c in actual_val if c not in ["：", "”", "’"])
                if has_zenkaku:
                    warnings.append(f"セル {cell_ref}: 数式または関数に全角文字が紛れ込んでいる可能性があります ('{actual_val}')")

            # 型チェック
            if expected_type == "NUMBER":
                try:
                    float(actual_val)
                except ValueError:
                    if not actual_val.startswith("="):
                        errors.append(f"セル {cell_ref}: 数値が期待されていますが '{actual_val}' が入力されています")
                        detail["status"] = "ERROR"
            elif expected_type == "DATE":
                if not is_date_cell(cell):
                    warnings.append(f"セル {cell_ref}: 日付形式が正しく設定されていない可能性があります")

            # ステージ固有ロジック
            if stage_id == 1:
                if actual_val.startswith("="):
                    errors.append(f"セル {cell_ref}: Stage 1では数式ではなくベタ打ちの値を入力してください")
                    detail["status"] = "ERROR"
            
            elif stage_id == 2:
                if not actual_val.startswith("="):
                    errors.append(f"セル {cell_ref}: 計算式(=)から始まる形式で入力してください")
                    detail["status"] = "ERROR"
                elif any(fn in actual_val.upper() for fn in ["SUM", "AVERAGE", "COUNT"]):
                    errors.append(f"セル {cell_ref}: 関数を使わずに基本四則演算(+, -, *, /)で記述してください")
                    detail["status"] = "ERROR"

            elif stage_id == 3:
                if expected_val and norm_actual != norm_expected:
                    if not "$" in actual_val:
                        warnings.append(f"セル {cell_ref}: 絶対参照($)が使われていない可能性があります")

            elif stage_id == 4:
                if "VLOOKUP" in actual_val.upper():
                    if "FALSE" not in actual_val.upper() and ",0)" not in actual_val:
                        warnings.append(f"セル {cell_ref}: VLOOKUPの第4引数(完全一致)が省略されているか非推奨の設定です")

            elif stage_id == 5:
                # 複合判定
                pass

            # エラー値チェック
            if actual_val in ["#REF!", "#VALUE!", "#N/A!", "#NAME?", "#DIV/0!"]:
                errors.append(f"セル {cell_ref}: エラー値 '{actual_val}' が検出されました")
                detail["status"] = "ERROR"

            # 期待値比較
            if expected_val and norm_actual != norm_expected and detail["status"] == "OK":
                errors.append(f"セル {cell_ref}: 期待値 '{expected_val}' と一致しません (入力値: '{actual_val}')")
                detail["status"] = "ERROR"
                detail["msg"] = f"期待値: {expected_val}"

            cell_details.append(detail)

        # Stage 6 グラフ判定
        if stage_id == 6:
            charts = ws._charts
            if not charts:
                errors.append("シート内にグラフが見つかりません。指示通りのグラフを作成してください。")
            else:
                for idx, chart in enumerate(charts):
                    title_text = "(未設定)"
                    if chart.title:
                        try:
                            title_text = chart.title.tx.rich.p[0].r[0].t
                        except Exception:
                            title_text = str(chart.title)
                    
                    if not chart.series:
                        warnings.append(f"グラフ{idx+1}: データが正しく割り当てられていない空のグラフの可能性があります")

        # Stage 7 ピボットテーブル存在確認 (Zip構造参照)
        if stage_id == 7:
            has_pivot = False
            for name in wb.sheetnames:
                # openpyxlでピボットテーブルのキャッシュがあるか判定
                ws_target = wb[name]
                if hasattr(ws_target, '_pivots') and len(ws_target._pivots) > 0:
                    has_pivot = True
                    break
            if not has_pivot:
                # ZIP構造として探す簡易判定
                import zipfile
                with zipfile.ZipFile(filepath, 'r') as z:
                    pivot_files = [f for f in z.namelist() if 'pivotTable' in f or 'pivotCache' in f]
                    if pivot_files:
                        has_pivot = True

            if not has_pivot:
                errors.append("ファイル内にピボットテーブルが検出されませんでした。")

    except Exception as e:
        errors.append(f"ファイル解析エラー: {str(e)}")

    return errors, warnings, cell_details

@app.route('/')
def index():
    url = f"{SUPABASE_URL}/stages?select=*&order=id.asc"
    resp = requests.get(url, headers=HEADERS)
    stages = resp.json() if resp.status_code == 200 else []
    return render_template('index.html', stages=stages)

@app.route('/upload', methods=['POST'])
def upload_file():
    student_name = request.form.get('student_name')
    stage_id = int(request.form.get('stage_id'))
    
    if 'excel_file' not in request.files:
        flash('ファイルが選択されていません', 'danger')
        return redirect(url_for('index'))
        
    file = request.files['excel_file']
    if file.filename == '' or not allowed_file(file.filename):
        flash('有効なExcelファイル(.xlsx, .xls)を選択してください', 'danger')
        return redirect(url_for('index'))

    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)

    errors, warnings, cell_details = analyze_excel_file(filepath, stage_id)

    if errors:
        status = "再提出が必要"
    elif warnings or stage_id in [6, 7]:
        status = "目視レビュー待ち"
    else:
        status = "合格"

    # Supabaseへ投稿登録
    payload = {
        "student_name": student_name,
        "stage_id": stage_id,
        "filename": filename,
        "status": status,
        "auto_feedback": "\n".join(errors + warnings) if (errors or warnings) else "自動チェックOK"
    }
    
    requests.post(f"{SUPABASE_URL}/submissions", json=payload, headers=HEADERS)

    return render_template('result.html', 
                           student_name=student_name, 
                           stage_id=stage_id, 
                           status=status, 
                           errors=errors, 
                           warnings=warnings,
                           cell_details=cell_details)

@app.route('/admin')
def admin():
    url = f"{SUPABASE_URL}/submissions?select=*,stages(title)&order=created_at.desc"
    resp = requests.get(url, headers=HEADERS)
    submissions = resp.json() if resp.status_code == 200 else []
    return render_template('admin.html', submissions=submissions)

@app.route('/admin/review', methods=['POST'])
def review():
    sub_id = request.form.get('submission_id')
    status = request.form.get('status')
    feedback = request.form.get('teacher_feedback')
    
    url = f"{SUPABASE_URL}/submissions?id=eq.{sub_id}"
    payload = {
        "status": status,
        "teacher_feedback": feedback
    }
    requests.patch(url, json=payload, headers=HEADERS)
    flash('レビューを保存しました', 'success')
    return redirect(url_for('admin'))

if __name__ == '__main__':
    app.run(debug=True, port=5000)